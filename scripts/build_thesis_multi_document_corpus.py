#!/usr/bin/env python3
"""Build a controlled multi-document thesis corpus and evaluation index."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

CORPUS_DIR = ROOT / "data/thesis_corpus"
DATASET_PATH = ROOT / "data/evaluation/thesis_multi_document_dataset.json"
CHROMA_DIR = ROOT / "data/chroma_thesis_eval"
BM25_PATH = ROOT / "data/bm25_thesis_eval.pkl"
GRAPH_PATH = ROOT / "data/graphs/thesis_multi_document_graph.pkl"


DOCS = [
    {
        "filename": "thesis_doc_01_rag_system_design.docx",
        "title": "Enterprise Knowledge Base RAG System Design",
        "sections": [
            (
                "Purpose and Operating Context",
                [
                    "This design document describes a retrieval augmented generation system for an enterprise knowledge base. The system is intended for staff who need grounded answers from policy documents, technical manuals, incident reports and deployment notes. The main requirement is not only fluent response generation, but traceable response generation. Every factual answer should be supported by retrieved source passages and should include inline citations that point back to the available context.",
                    "The system is designed for document grounded question answering. It does not train a new language model. Instead, it combines document ingestion, hierarchical chunking, vector retrieval, keyword retrieval, answer synthesis and reliability checking around an existing language model. The design assumes that uploaded documents are the primary source of truth.",
                ],
            ),
            (
                "Ingestion and Retrieval Design",
                [
                    "Documents are loaded from PDF, DOCX and text files. Each document is split into large parent chunks and smaller child chunks. Parent chunks preserve broader context, while child chunks provide precise searchable units. This hierarchy allows the retriever to search smaller units while still returning enough surrounding evidence for answer generation.",
                    "Vector retrieval is used for semantic matching. It is useful when a user asks a paraphrased question, such as asking about response reliability when the document uses the phrase answer grounding. BM25 keyword retrieval is used for exact terms, names, measurements and policy labels. The retrieval coordinator combines vector and BM25 results so that semantic evidence and exact lexical evidence can both be considered.",
                ],
            ),
            (
                "Citation and Reliability Policy",
                [
                    "The writer must cite evidence using numbered inline citations such as [1] or [2]. A citation is valid only when the cited retrieved chunk supports the sentence that uses it. The writer should not cite a chunk only because it is generally related to the topic. If the retrieved context does not contain the answer, the writer must state that the provided documents do not contain the requested information.",
                    "The reliability gate checks whether the answer contains citations, whether citation numbers refer to available chunks and whether retrieved context exists. If the gate fails but context is available, the system can fall back to a deterministic evidence summary. This prevents unsupported answers from being returned as confident final responses.",
                ],
            ),
            (
                "Known Failure Cases",
                [
                    "A vector only baseline can fail when a question contains exact terms that are rare in the document. It can also retrieve semantically related passages that do not actually answer the question. Another failure case occurs when a broad summary question needs evidence from multiple parts of a document but only one narrow chunk is returned.",
                    "The agentic system reduces these risks by coordinating retrieval methods, synthesising duplicate evidence and checking the final answer. The tradeoff is latency. The baseline is simpler and faster because it performs one vector retrieval step and one generation step. The agentic system is slower because it performs multiple retrieval and quality control steps.",
                ],
            ),
        ],
    },
    {
        "filename": "thesis_doc_02_multi_agent_workflow.docx",
        "title": "Hierarchical Multi Agent Workflow Specification",
        "sections": [
            (
                "Architecture Overview",
                [
                    "The proposed workflow uses a three level hierarchy. The strategic level contains the Planner Agent. The tactical level contains the Query Decomposer, Retrieval Coordinator, Validator, Synthesis Agent, Writer Agent, Critic Agent and Reliability Gate. The operational level contains the Vector Search Agent, Keyword Search Agent and optional Graph Search Agent.",
                    "The hierarchy separates decision making from execution. The Planner Agent estimates query complexity and selects a strategy. Simple questions can use direct retrieval. Complex questions can use decomposition and multi step retrieval. Relationship questions can activate graph based reasoning when a knowledge graph is available.",
                ],
            ),
            (
                "Agent Responsibilities",
                [
                    "The Query Decomposer breaks complex questions into smaller sub questions. The Retrieval Coordinator launches the retrieval swarm and gathers evidence from vector, keyword and graph retrieval agents. The Validator checks whether the retrieved evidence is sufficient before the system moves to answer generation.",
                    "The Synthesis Agent removes duplicate chunks and ranks evidence before writing. The Writer Agent creates the final answer with inline citations. The Critic Agent reviews answer quality and can request regeneration when the answer is incomplete, weakly grounded or poorly cited. The Reliability Gate performs deterministic final checks before the response is returned.",
                ],
            ),
            (
                "Self Reflection Loop",
                [
                    "Self reflection is implemented through the Critic Agent and regeneration loop. After the Writer Agent produces an answer, the Critic Agent evaluates accuracy, relevance, completeness and citation quality. If the answer does not satisfy the threshold, the Critic Agent returns feedback and the Writer Agent generates a revised answer.",
                    "The reflection loop is limited by a maximum number of iterations. This prevents the workflow from becoming too slow or entering an endless cycle. In the benchmark setting, the maximum number of critic iterations can be reduced to keep experiments practical.",
                ],
            ),
            (
                "Workflow Trace",
                [
                    "A typical complex query follows this route: Planner Agent, Query Decomposer, Retrieval Coordinator, Validator, Synthesis Agent, Writer Agent, Critic Agent and Reliability Gate. The Retrieval Coordinator may call the Vector Search Agent for semantic evidence, the Keyword Search Agent for exact evidence and the Graph Search Agent for entity relationship evidence.",
                    "The system records metadata such as strategy, retrieval round, retrieved chunk count, validation status, critic score and reliability gate result. These metadata fields support evaluation because they make the workflow more transparent than a single black box generation call.",
                ],
            ),
        ],
    },
    {
        "filename": "thesis_doc_03_graph_reasoning_case_study.docx",
        "title": "Graph Reasoning Case Study for Document Grounded QA",
        "sections": [
            (
                "Case Background",
                [
                    "A support organisation maintains documents about the Orion deployment incident. The incident involved the Orion Service, the Billing API, the Cache Layer, the Search Index and the Audit Dashboard. The main problem was delayed invoice visibility after a schema migration changed the event payload used by downstream services.",
                    "The Orion Service publishes invoice events to the Billing API. The Billing API writes normalised invoice records into the Cache Layer. The Search Index reads from the Cache Layer every fifteen minutes. The Audit Dashboard displays invoice status using data from the Search Index. This chain means that a failure in the Cache Layer can affect the Audit Dashboard even when the dashboard itself is healthy.",
                ],
            ),
            (
                "Entities and Relationships",
                [
                    "The schema migration introduced a field named invoice_state, replacing the older field named status_code. The Billing API accepted the new field, but the Cache Layer still expected status_code. As a result, records were stored without a searchable status value. The Search Index then indexed incomplete records, and the Audit Dashboard showed invoices as pending review.",
                    "The Platform Team owns the Orion Service and the Search Index. The Finance Systems Team owns the Billing API. The Reliability Team owns the Cache Layer and monitors the recovery procedure. The Compliance Team uses the Audit Dashboard for month end evidence review.",
                ],
            ),
            (
                "Recovery Procedure",
                [
                    "The recovery procedure has four steps. First, the Finance Systems Team maps invoice_state back to status_code for affected records. Second, the Reliability Team clears stale Cache Layer entries. Third, the Platform Team triggers a Search Index rebuild. Fourth, the Compliance Team verifies the Audit Dashboard against a sample of corrected invoices.",
                    "The root cause is not the Audit Dashboard. The dashboard only displays the indexed state. The direct cause is the mismatch between invoice_state and status_code in the Cache Layer. The downstream effect is incomplete indexing, which then affects compliance review.",
                ],
            ),
            (
                "Risk and Evidence Notes",
                [
                    "The highest risk is repeating the incident during future schema migrations. The recommended control is a compatibility contract between the Billing API and the Cache Layer. Another control is an automated relationship test that checks whether the Search Index and Audit Dashboard receive the expected status value after migration.",
                    "The document does not provide customer names, exact revenue impact, the number of affected invoices or a final regulatory conclusion. Any answer about those details should state that the information is missing from the provided document.",
                ],
            ),
        ],
    },
    {
        "filename": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "title": "RAG Evaluation and Ablation Report",
        "sections": [
            (
                "Evaluation Setup",
                [
                    "This report evaluates a baseline vector only RAG system and a proposed hierarchical agentic RAG system. The baseline uses query embedding, top five vector retrieval and one writer call. The proposed system uses hybrid retrieval, evidence synthesis, self reflection and a reliability gate. Both systems are evaluated on the same document grounded questions.",
                    "The evaluation uses RAGAS metrics and custom reliability metrics. RAGAS metrics include faithfulness, answer relevancy, context precision and context recall. Custom metrics include citation rate, citation based context usage, average answer length, average latency and retrieved chunk count.",
                ],
            ),
            (
                "Main Result Summary",
                [
                    "In the controlled benchmark, the baseline achieved an average RAGAS overall score of 0.742. The proposed agentic RAG system achieved an average RAGAS overall score of 0.846. The largest improvement was observed in faithfulness, where the baseline scored 0.710 and the proposed system scored 0.890.",
                    "The citation rate improved from 0.78 for the baseline to 0.96 for the proposed system. Citation based context usage improved from 0.46 to 0.61. Average latency increased from 18.4 seconds to 31.7 seconds. This indicates that the proposed system improves reliability and evidence use at the cost of slower response time.",
                ],
            ),
            (
                "Ablation Study",
                [
                    "The ablation study tested three component removals. Removing BM25 reduced exact fact retrieval performance by 12 percent. Removing the Critic Agent reduced faithfulness by 9 percent because weakly cited answers were no longer regenerated. Removing the Reliability Gate increased unsupported final answers from 3 percent to 11 percent.",
                    "Graph search was most useful for relationship questions. In the graph reasoning subset, the proposed system with graph search scored 0.83 overall, while the same system without graph search scored 0.71. The improvement was linked to questions that required connecting services, teams, risks and downstream effects.",
                ],
            ),
            (
                "Limitations",
                [
                    "The benchmark is a controlled document benchmark, not a large public leaderboard. The number of documents is small, but the documents are selected to cover different reasoning needs. The RAGAS judge can also be sensitive to wording when the correct answer is an honest statement that information is missing.",
                    "The report does not claim that agentic RAG is always better than baseline RAG. For simple factual questions, the baseline can be faster and sometimes sufficient. The proposed system is most useful when the question requires citations, comparison, missing information handling, relationship reasoning or quality control.",
                ],
            ),
        ],
    },
]


DATASET = [
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "What is the main purpose of the enterprise knowledge base RAG system?",
        "ground_truth": "The system is designed for document grounded question answering over enterprise documents, with traceable answers supported by retrieved passages and inline citations.",
        "category": "summary",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "Which file formats are supported during document ingestion?",
        "ground_truth": "The document states that PDF, DOCX and text files are supported during ingestion.",
        "category": "fact extraction",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "How do parent chunks and child chunks work in the retrieval design?",
        "ground_truth": "Parent chunks preserve broader context, while child chunks provide precise searchable units. The retriever searches smaller child chunks while returning enough surrounding evidence for answer generation.",
        "category": "method explanation",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "Compare vector retrieval and BM25 keyword retrieval in this design.",
        "ground_truth": "Vector retrieval supports semantic matching for paraphrased questions, while BM25 keyword retrieval supports exact terms, names, measurements and policy labels. The coordinator combines both to improve evidence coverage.",
        "category": "comparison",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "How does the reliability gate relate citations to retrieved chunks?",
        "ground_truth": "The reliability gate checks whether the answer contains citations, whether citation numbers refer to available chunks and whether retrieved context exists.",
        "category": "relationship reasoning",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "Which evidence best supports the claim that the proposed system is more reliable than vector only RAG?",
        "ground_truth": "The document states that the agentic system coordinates retrieval methods, synthesises duplicate evidence and checks the final answer, while vector only RAG can miss exact terms or return related but insufficient passages.",
        "category": "evidence selection",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "What tradeoff is introduced by using the agentic system instead of the baseline?",
        "ground_truth": "The agentic system reduces reliability risks but increases latency because it performs multiple retrieval and quality control steps, while the baseline is faster and simpler.",
        "category": "limitations",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_01_rag_system_design.docx",
        "question": "What exact production cost is reported for operating the enterprise RAG system?",
        "ground_truth": "The document does not provide an exact production cost for operating the enterprise RAG system.",
        "category": "missing information",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "What are the three levels in the hierarchical multi agent workflow?",
        "ground_truth": "The three levels are the strategic level with the Planner Agent, the tactical level with agents such as the Decomposer, Coordinator, Validator, Writer, Critic and Reliability Gate, and the operational level with Vector, Keyword and optional Graph Search Agents.",
        "category": "summary",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "Which agents are listed at the operational level?",
        "ground_truth": "The operational level contains the Vector Search Agent, Keyword Search Agent and optional Graph Search Agent.",
        "category": "fact extraction",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "How does the self reflection loop work?",
        "ground_truth": "After the Writer Agent produces an answer, the Critic Agent evaluates accuracy, relevance, completeness and citation quality. If the answer is below threshold, feedback is returned and the Writer Agent generates a revised answer.",
        "category": "method explanation",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "Compare the Planner Agent and the Retrieval Coordinator.",
        "ground_truth": "The Planner Agent estimates query complexity and selects a strategy, while the Retrieval Coordinator launches the retrieval swarm and gathers evidence from vector, keyword and graph retrieval agents.",
        "category": "comparison",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "How is the Critic Agent connected to the Writer Agent in the workflow?",
        "ground_truth": "The Critic Agent reviews the Writer Agent output and can request regeneration, sending feedback so the Writer Agent can create a revised answer.",
        "category": "relationship reasoning",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "Which evidence shows that the workflow is transparent for evaluation?",
        "ground_truth": "The document states that the system records metadata such as strategy, retrieval round, retrieved chunk count, validation status, critic score and reliability gate result.",
        "category": "evidence selection",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "What limitation is placed on the reflection loop?",
        "ground_truth": "The reflection loop is limited by a maximum number of iterations to prevent the workflow from becoming too slow or entering an endless cycle.",
        "category": "limitations",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_02_multi_agent_workflow.docx",
        "question": "What exact hardware configuration is used by the Planner Agent?",
        "ground_truth": "The document does not provide an exact hardware configuration for the Planner Agent.",
        "category": "missing information",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "What incident is described in the graph reasoning case study?",
        "ground_truth": "The document describes the Orion deployment incident, where delayed invoice visibility followed a schema migration that changed the event payload used by downstream services.",
        "category": "summary",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "Which teams own the Billing API and the Cache Layer?",
        "ground_truth": "The Finance Systems Team owns the Billing API, and the Reliability Team owns the Cache Layer.",
        "category": "fact extraction",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "What are the four steps in the recovery procedure?",
        "ground_truth": "The four steps are mapping invoice_state back to status_code, clearing stale Cache Layer entries, triggering a Search Index rebuild and verifying the Audit Dashboard against corrected invoices.",
        "category": "method explanation",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "Compare the direct cause and the downstream effect in the Orion incident.",
        "ground_truth": "The direct cause is the mismatch between invoice_state and status_code in the Cache Layer. The downstream effect is incomplete indexing, which affects the Audit Dashboard and compliance review.",
        "category": "comparison",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "How does the Billing API affect the Audit Dashboard through the Cache Layer and Search Index?",
        "ground_truth": "The Billing API writes normalised invoice records into the Cache Layer. The Search Index reads from the Cache Layer, and the Audit Dashboard displays invoice status using data from the Search Index.",
        "category": "relationship reasoning",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "Which evidence shows that the Audit Dashboard was not the root cause?",
        "ground_truth": "The document states that the dashboard only displays the indexed state and that the direct cause is the mismatch between invoice_state and status_code in the Cache Layer.",
        "category": "evidence selection",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "What is the highest risk identified for future schema migrations?",
        "ground_truth": "The highest risk is repeating the incident during future schema migrations.",
        "category": "limitations",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_03_graph_reasoning_case_study.docx",
        "question": "How many invoices were affected by the Orion incident?",
        "ground_truth": "The document does not provide the number of affected invoices.",
        "category": "missing information",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "What systems are compared in the evaluation report?",
        "ground_truth": "The report compares a baseline vector only RAG system with a proposed hierarchical agentic RAG system.",
        "category": "summary",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "What average RAGAS overall scores are reported for baseline and proposed systems?",
        "ground_truth": "The baseline achieved an average RAGAS overall score of 0.742, while the proposed agentic RAG system achieved 0.846.",
        "category": "fact extraction",
        "difficulty": "easy",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "How are RAGAS metrics and custom reliability metrics used in the report?",
        "ground_truth": "RAGAS metrics measure faithfulness, answer relevancy, context precision and context recall. Custom metrics measure citation rate, citation based context usage, answer length, latency and retrieved chunk count.",
        "category": "method explanation",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "Compare the reliability improvement and latency cost of the proposed system.",
        "ground_truth": "The proposed system improves reliability, with citation rate increasing from 0.78 to 0.96 and context usage from 0.46 to 0.61, but average latency increases from 18.4 seconds to 31.7 seconds.",
        "category": "comparison",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "How does graph search relate to the graph reasoning subset?",
        "ground_truth": "Graph search is most useful for relationship questions. In the graph reasoning subset, the system with graph search scored 0.83 overall, while the same system without graph search scored 0.71.",
        "category": "relationship reasoning",
        "difficulty": "hard",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "Which evidence supports the claim that removing the Reliability Gate is risky?",
        "ground_truth": "The ablation study states that removing the Reliability Gate increased unsupported final answers from 3 percent to 11 percent.",
        "category": "evidence selection",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "What limitation does the report state about the benchmark?",
        "ground_truth": "The benchmark is a controlled document benchmark rather than a large public leaderboard, and RAGAS can be sensitive to wording for honest missing information answers.",
        "category": "limitations",
        "difficulty": "medium",
    },
    {
        "document": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "question": "What public leaderboard ranking did the proposed agentic RAG system achieve?",
        "ground_truth": "The document does not provide any public leaderboard ranking for the proposed agentic RAG system.",
        "category": "missing information",
        "difficulty": "hard",
    },
]


GRAPH_RELATIONSHIPS = [
    ("planner agent", "selects", "strategy"),
    ("retrieval coordinator", "calls", "vector search agent"),
    ("retrieval coordinator", "calls", "keyword search agent"),
    ("retrieval coordinator", "calls", "graph search agent"),
    ("writer agent", "is_reviewed_by", "critic agent"),
    ("critic agent", "requests", "regeneration"),
    ("reliability gate", "checks", "citations"),
    ("billing api", "writes_to", "cache layer"),
    ("cache layer", "feeds", "search index"),
    ("search index", "feeds", "audit dashboard"),
    ("schema migration", "changed", "invoice_state"),
    ("cache layer", "expected", "status_code"),
    ("graph search", "improves", "relationship questions"),
    ("bm25", "improves", "exact fact retrieval"),
    ("critic agent", "improves", "faithfulness"),
]


def _set_styles(doc: WordDocument) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Heading 1", 16, "1F4D78"),
        ("Heading 2", 13, "2E74B5"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def write_docx_files() -> None:
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    for spec in DOCS:
        doc = WordDocument()
        _set_styles(doc)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(spec["title"])
        run.bold = True
        run.font.size = Pt(18)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor.from_string("1F4D78")

        for heading, paragraphs in spec["sections"]:
            doc.add_heading(heading, level=2)
            for text in paragraphs:
                p = doc.add_paragraph(text)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        path = CORPUS_DIR / spec["filename"]
        doc.core_properties.author = "Huang Xuan"
        doc.core_properties.title = spec["title"]
        doc.core_properties.subject = "Controlled thesis evaluation corpus"
        doc.save(path)
        print(f"Wrote {path}")


def write_dataset() -> None:
    DATASET_PATH.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "description": (
            "Controlled multi-document thesis benchmark for hierarchical "
            "agentic RAG evaluation."
        ),
        "documents": [doc["filename"] for doc in DOCS],
        "test_cases": DATASET,
    }
    DATASET_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Wrote {DATASET_PATH}")


def _slug(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_").lower()
    return slug[:40] or "doc"


def _prefix_chunk_ids(parent_chunks, child_chunks, prefix: str) -> None:
    parent_id_map = {}
    for parent in parent_chunks:
        old_id = parent.chunk_id
        new_id = f"{prefix}_{old_id}"
        parent_id_map[old_id] = new_id
        parent.chunk_id = new_id

    for child in child_chunks:
        old_id = child.chunk_id
        child.chunk_id = f"{prefix}_{old_id}"
        if child.parent_id in parent_id_map:
            child.parent_id = parent_id_map[child.parent_id]

    for parent in parent_chunks:
        parent.children_ids = [
            f"{prefix}_{child_id}" for child_id in parent.children_ids
        ]


def build_index() -> None:
    from dotenv import load_dotenv
    from src.ingestion.document_loader import DocumentLoader
    from src.ingestion.embedder import EmbeddingGenerator
    from src.ingestion.hierarchical_chunker import HierarchicalChunker
    from src.storage.chroma_store import ChromaVectorStore
    from src.retrieval.bm25_index import BM25Index

    load_dotenv(ROOT / ".env")

    if CHROMA_DIR.exists():
        shutil.rmtree(CHROMA_DIR)
    if BM25_PATH.exists():
        BM25_PATH.unlink()

    store = ChromaVectorStore(persist_directory=str(CHROMA_DIR))
    embedder = EmbeddingGenerator()
    loader = DocumentLoader()
    chunker = HierarchicalChunker(parent_size=1200, child_size=350, child_overlap=50)

    for spec in DOCS:
        path = CORPUS_DIR / spec["filename"]
        loaded = loader.load(str(path))
        metadata = {
            "filename": spec["filename"],
            "source": "thesis_controlled_corpus",
            "title": spec["title"],
        }
        parent_chunks, child_chunks = chunker.chunk_text(
            loaded.text,
            doc_id=loaded.doc_id,
            metadata=metadata,
        )
        _prefix_chunk_ids(parent_chunks, child_chunks, _slug(path.stem))

        for chunk in parent_chunks + child_chunks:
            chunk.embedding = embedder.generate([chunk.text])[0]

        store.add_chunks(parent_chunks, child_chunks, filename=spec["filename"])
        print(
            f"Indexed {spec['filename']}: "
            f"{len(parent_chunks)} parents, {len(child_chunks)} children"
        )

    bm25 = BM25Index(index_path=str(BM25_PATH))
    bm25.build_from_vector_store(store)
    bm25.save()
    print(f"Wrote {BM25_PATH}")


def build_graph() -> None:
    from src.graph.graph_builder import KnowledgeGraph
    from src.graph.relationship_extractor import Relationship

    kg = KnowledgeGraph()
    for source, relation, target in GRAPH_RELATIONSHIPS:
        kg.add_relationship(Relationship(source, relation, target, confidence=0.95))
    kg.save(str(GRAPH_PATH))
    print(f"Wrote {GRAPH_PATH}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-index", action="store_true")
    parser.add_argument("--no-graph", action="store_true")
    args = parser.parse_args()

    write_docx_files()
    write_dataset()
    if not args.no_index:
        build_index()
    if not args.no_graph:
        build_graph()


if __name__ == "__main__":
    main()
