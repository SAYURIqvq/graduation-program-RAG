#!/usr/bin/env python3
"""Build the complex thesis stress benchmark corpus, dataset, indexes, and graph."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

CORPUS_DIR = ROOT / "data/thesis_corpus"
COMPLEX_DATASET_PATH = ROOT / "data/evaluation/thesis_complex_reasoning_dataset.json"
CHROMA_DIR = ROOT / "data/chroma_thesis_eval"
BM25_PATH = ROOT / "data/bm25_thesis_eval.pkl"
GRAPH_PATH = ROOT / "data/graphs/thesis_multi_document_graph.pkl"


DOCS = [
    {
        "filename": "thesis_doc_01_rag_system_design.docx",
        "title": "Enterprise Knowledge Base RAG System Design",
        "sections": [
            (
                "Purpose and Evidence Policy",
                [
                    "This design document describes an enterprise retrieval augmented generation system for policy documents, technical manuals, incident reports and deployment notes. The main goal is not only fluent response generation, but traceable response generation. Every factual answer should be grounded in retrieved source passages and should include inline citations that point to available evidence.",
                    "The system does not train a new language model. Uploaded documents are treated as the source of truth. The RAG pipeline uses hierarchical chunking, vector retrieval, BM25 keyword retrieval, synthesis and reliability checks around an existing language model.",
                ],
            ),
            (
                "Chunking and Retrieval Flow",
                [
                    "Documents are split into large parent chunks and smaller child chunks. Child chunks are searched because they are precise. Parent chunks are returned because they preserve enough surrounding context for answer generation. This design prevents the writer from seeing only a narrow fragment when a question requires a broader explanation.",
                    "Vector retrieval is useful for paraphrased questions. BM25 is useful for exact terms, names, rare policy labels and measurements. The retrieval coordinator merges vector and BM25 results, then removes duplicates before the writer receives the evidence.",
                ],
            ),
            (
                "Citation Rules and Fallback",
                [
                    "The writer must use numbered inline citations such as [1] or [2]. A citation is valid only when the cited chunk directly supports the sentence. The writer should not cite a chunk merely because it is generally related to the topic.",
                    "If the retrieved context does not contain the answer, the writer must state that the provided documents do not contain the requested information. If the reliability gate fails but retrieved context exists, the system can fall back to a deterministic evidence summary instead of returning an unsupported fluent answer.",
                ],
            ),
            (
                "Baseline Failure Modes",
                [
                    "A vector only baseline can fail when a question contains exact terms that are rare in the document. It can also retrieve a semantically related passage that does not answer the question. Broad summary questions can fail when only one narrow child chunk is returned and important policy details appear in a different section.",
                    "The agentic system reduces these risks by coordinating retrieval methods, synthesising duplicate evidence and checking the final answer. The tradeoff is latency. The baseline is faster because it performs one vector retrieval step and one generation step. The agentic system is slower because it performs multiple retrieval and quality control steps.",
                ],
            ),
            (
                "Reliability Example",
                [
                    "When a user asks whether a deployment policy allows unapproved rollback, the answer must combine the rollback rule, the citation rule and the missing information rule. If the document contains rollback requirements but does not contain an approval exception, the correct answer should describe the available rollback rule and explicitly state that an approval exception is not provided.",
                    "The evidence support policy therefore connects three ideas: retrieved chunks provide facts, citations connect claims to chunks and the reliability gate prevents unsupported claims from becoming final answers.",
                ],
            ),
        ],
    },
    {
        "filename": "thesis_doc_02_multi_agent_workflow.docx",
        "title": "Hierarchical Multi Agent Workflow Specification",
        "sections": [
            (
                "Three Level Hierarchy",
                [
                    "The proposed workflow uses a three level hierarchy. The strategic level contains the Planner Agent. The tactical level contains the Query Decomposer, Retrieval Coordinator, Validator, Synthesis Agent, Writer Agent, Critic Agent and Reliability Gate. The operational level contains the Vector Search Agent, Keyword Search Agent and Graph Search Agent.",
                    "The hierarchy separates decision making from execution. The Planner estimates query complexity and selects a strategy. Simple questions can use direct retrieval. Complex questions use decomposition and multi step retrieval. Relationship questions activate graph based reasoning when a knowledge graph is available.",
                ],
            ),
            (
                "Agent Responsibilities",
                [
                    "The Query Decomposer breaks complex questions into smaller sub questions. The Retrieval Coordinator launches the retrieval swarm and gathers evidence from vector, keyword and graph retrieval agents. The Validator checks whether retrieved evidence is sufficient before the system moves to answer generation.",
                    "The Synthesis Agent removes duplicate chunks and ranks evidence. The Writer Agent creates the final answer with inline citations. The Critic Agent reviews answer quality and can request regeneration when the answer is incomplete, weakly grounded or poorly cited. The Reliability Gate performs deterministic final checks before the response is returned.",
                ],
            ),
            (
                "Self Reflection and Regeneration",
                [
                    "Self reflection is implemented through the Critic Agent and regeneration loop. After the Writer Agent produces an answer, the Critic Agent evaluates accuracy, relevance, completeness and citation quality. If the answer does not satisfy the threshold, the Critic Agent returns feedback and the Writer Agent generates a revised answer.",
                    "The loop is limited by a maximum number of iterations to prevent excessive latency or endless revision. During benchmark measurement, critic regeneration can be disabled while critic scores are still recorded. This makes the experiment practical while preserving evidence of self reflection behaviour.",
                ],
            ),
            (
                "Validation and Retrieval Retry",
                [
                    "The Validator can decide that retrieved evidence is insufficient. In that case, the workflow returns to the Retrieval Coordinator for another retrieval round. This retry mechanism is important when the first retrieval round finds a related passage but misses a necessary supporting detail.",
                    "A typical complex query follows this route: Planner, Query Decomposer, Retrieval Coordinator, Validator, Synthesis Agent, Writer Agent, Critic Agent and Reliability Gate. The route is longer than baseline RAG, but it records strategy, retrieval round, validation status, critic score and reliability gate result.",
                ],
            ),
            (
                "Transparent Evaluation",
                [
                    "The workflow stores metadata such as selected strategy, sub query count, retrieval agents used, retrieved chunk count, validation decision, critic score and final reliability status. These fields make the process more transparent than a single black box generation call.",
                    "The workflow is not intended to replace baseline RAG for every query. For direct facts, baseline retrieval can be faster and sufficient. The full hierarchy is intended for high reliability questions that need decomposition, multi source evidence, missing information handling or relationship reasoning.",
                ],
            ),
        ],
    },
    {
        "filename": "thesis_doc_03_graph_reasoning_case_study.docx",
        "title": "Graph Reasoning Case Study for Document Grounded QA",
        "sections": [
            (
                "Incident Background",
                [
                    "A support organisation maintains documents about the Orion deployment incident. The incident involved the Orion Service, Billing API, Cache Layer, Search Index, Audit Dashboard and Compliance Review process. The problem was delayed invoice visibility after a schema migration changed the event payload used by downstream services.",
                    "The Orion Service publishes invoice events to the Billing API. The Billing API writes normalised invoice records into the Cache Layer. The Search Index reads from the Cache Layer every fifteen minutes. The Audit Dashboard displays invoice status using the Search Index. Compliance Review depends on the Audit Dashboard for month end evidence.",
                ],
            ),
            (
                "Schema Mismatch Chain",
                [
                    "The schema migration introduced invoice_state and replaced the older field status_code. The Billing API accepted invoice_state, but the Cache Layer still expected status_code. Records were therefore stored without a searchable status value. The Search Index indexed incomplete records, and the Audit Dashboard showed invoices as pending review.",
                    "The dashboard was not the root cause. It only displayed the indexed state. The direct cause was the mismatch between invoice_state and status_code in the Cache Layer. The downstream effect was incomplete indexing, which then affected dashboard visibility and Compliance Review.",
                ],
            ),
            (
                "Owners and Dependencies",
                [
                    "The Platform Team owns the Orion Service and Search Index. The Finance Systems Team owns the Billing API. The Reliability Team owns the Cache Layer and monitors recovery. The Compliance Team uses the Audit Dashboard for evidence review.",
                    "The dependency path is Orion Service to Billing API to Cache Layer to Search Index to Audit Dashboard to Compliance Review. A failure in the Cache Layer can therefore affect Compliance Review even if both the dashboard and the review process are healthy.",
                ],
            ),
            (
                "Recovery and Controls",
                [
                    "The recovery procedure has four steps. First, Finance Systems maps invoice_state back to status_code for affected records. Second, Reliability clears stale Cache Layer entries. Third, Platform triggers a Search Index rebuild. Fourth, Compliance verifies the Audit Dashboard against corrected invoices.",
                    "The highest future risk is repeating the incident during another schema migration. The recommended control is a compatibility contract between Billing API and Cache Layer. A second control is an automated relationship test that checks whether Search Index and Audit Dashboard receive the expected status value after migration.",
                ],
            ),
            (
                "Missing Evidence Boundary",
                [
                    "The document does not provide customer names, exact revenue impact, the number of affected invoices, a public incident severity grade or a final regulatory conclusion. Any answer about those details should state that the information is missing from the provided document.",
                    "The case is designed for graph based reasoning because correct answers must connect services, fields, teams, risks and downstream effects rather than retrieve a single isolated fact.",
                ],
            ),
        ],
    },
    {
        "filename": "thesis_doc_04_evaluation_and_ablation_report.docx",
        "title": "RAG Evaluation and Ablation Report",
        "sections": [
            (
                "Evaluation Setup",
                [
                    "This report evaluates a baseline vector only RAG system and a proposed hierarchical agentic RAG system. The baseline uses query embedding, top five vector retrieval and one writer call. The proposed system uses hybrid retrieval, evidence synthesis, self reflection and a reliability gate. Both systems are evaluated on the same document grounded questions.",
                    "RAGAS metrics include faithfulness, answer relevancy, context precision and context recall. Custom metrics include citation rate, citation based context usage, missing information accuracy, multi hop coverage, graph reasoning success, average answer length, latency and retrieved chunk count.",
                ],
            ),
            (
                "Main Result Summary",
                [
                    "In the controlled benchmark, the baseline achieved an average RAGAS overall score of 0.742. The proposed agentic RAG system achieved an average RAGAS overall score of 0.846. The largest improvement was faithfulness, where the baseline scored 0.710 and the proposed system scored 0.890.",
                    "Citation rate improved from 0.78 for the baseline to 0.96 for the proposed system. Citation based context usage improved from 0.46 to 0.61. Average latency increased from 18.4 seconds to 31.7 seconds. This indicates that the proposed system improves reliability and evidence use at the cost of slower response time.",
                ],
            ),
            (
                "Ablation Study",
                [
                    "The ablation study tested component removals. Removing BM25 reduced exact fact retrieval performance by 12 percent. Removing the Critic Agent reduced faithfulness by 9 percent because weakly cited answers were no longer regenerated. Removing the Reliability Gate increased unsupported final answers from 3 percent to 11 percent.",
                    "Graph search was most useful for relationship questions. In the graph reasoning subset, the proposed system with graph search scored 0.83 overall, while the same system without graph search scored 0.71. The improvement was linked to questions that required connecting services, teams, risks and downstream effects.",
                ],
            ),
            (
                "Latency and Scope",
                [
                    "The report does not claim that agentic RAG is always better than baseline RAG. For simple factual questions, baseline can be faster and sufficient. The proposed system is most useful when the question requires citations, comparison, missing information handling, relationship reasoning or quality control.",
                    "The benchmark is a controlled document benchmark, not a large public leaderboard. The number of documents is small, but the documents are selected to cover different reasoning needs. RAGAS can also be sensitive to wording when the correct answer is an honest statement that information is missing.",
                ],
            ),
            (
                "Recommended Interpretation",
                [
                    "The main interpretation is that baseline RAG is competitive for simple fact extraction, but hierarchical agentic RAG is stronger on complex reasoning and reliability oriented tasks. The evidence for this claim comes from graph reasoning performance, missing information behaviour, ablation results and validation retry traces.",
                    "The latency cost should be reported as a limitation rather than hidden. In high reliability settings, the additional latency is acceptable when the question needs evidence validation or multi hop reasoning. In low risk direct lookup settings, a baseline route may be preferable.",
                ],
            ),
        ],
    },
]


STRESS_APPENDIX_SECTIONS = {
    "thesis_doc_01_rag_system_design.docx": [
        (
            "Retrieval Stress Notes",
            [
                "A semantic retrieval note records that paraphrased questions often mention evidence, support, answer quality and reliability without naming the exact policy section. In those cases, vector retrieval can find a related chunk, but it may not include the citation rule, fallback rule and missing information rule in the same passage.",
                "A lexical retrieval note records that rare terms such as approval exception, status label, rollback window and deterministic evidence summary should be preserved during indexing. These terms are intentionally distributed across different parts of the document so a broad query must combine several pieces of evidence.",
            ],
        ),
        (
            "Citation Audit Appendix",
            [
                "The citation audit distinguishes direct support from topical similarity. Direct support means the cited passage states the fact needed by the sentence. Topical similarity means the passage talks about the same general topic but does not prove the exact claim.",
                "A citation is rejected when it points to a chunk that only discusses retrieval in general while the answer sentence claims a specific fallback rule, missing information rule or rollback exception. This is why the reliability gate checks citation numbers and available context together.",
            ],
        ),
        (
            "Rollback Evidence Appendix",
            [
                "The rollback evidence record says emergency rollback requires a deployment note, an affected service name and a cited operational rule. It does not list an approved exception for bypassing those requirements.",
                "When a question asks for an unapproved rollback exception, the answer should not invent an exception. It should describe the available rollback requirement and state that the exception is not contained in the provided documents.",
            ],
        ),
        (
            "Latency Tradeoff Appendix",
            [
                "The fast route performs one vector search and one writer call. It is appropriate for low risk direct lookup questions where the required fact is contained in one retrieved chunk.",
                "The high reliability route performs planning, decomposition, hybrid retrieval, validation, synthesis, writing, criticism and final gate checks. It is appropriate when evidence is scattered, missing or relationship based.",
            ],
        ),
        (
            "Decoy Retrieval Notes",
            [
                "Some retrieval notes mention citations, reliability and answer support only as general evaluation vocabulary. These notes are useful for testing whether a system can avoid treating broad topical matches as sufficient evidence.",
                "A broad passage about reliable RAG may be relevant to the question but still insufficient if it does not mention fallback behaviour, approval exceptions or direct citation support.",
            ],
        ),
    ],
    "thesis_doc_02_multi_agent_workflow.docx": [
        (
            "Planner Trace Appendix",
            [
                "The Planner records whether the query is simple, comparative, multi hop or relationship oriented. A simple label can skip decomposition, while a multi hop label passes the query to the Query Decomposer.",
                "The Planner does not itself answer the question. It selects the route that determines whether direct retrieval, decomposition or graph based retrieval should be attempted.",
            ],
        ),
        (
            "Validator Retry Appendix",
            [
                "The Validator checks whether retrieved evidence is sufficient for the sub questions created by the decomposer. If a related passage is present but a necessary supporting detail is missing, the Validator can request another retrieval round.",
                "A retry is not a guarantee that new evidence exists. It is a reliability behaviour showing that the system has detected a possible evidence gap before writing the final answer.",
            ],
        ),
        (
            "Critic Trace Appendix",
            [
                "The Critic Agent reviews the answer after writing. It scores accuracy, relevance, completeness and citation quality, then records whether the answer should be approved or regenerated.",
                "During time limited benchmarks, regeneration can be disabled while the critic decision and critic score are still stored. This setting preserves trace evidence of self reflection without making the evaluation impractically slow.",
            ],
        ),
        (
            "Transparency Appendix",
            [
                "The workflow trace stores selected strategy, sub query count, retrieval agent names, retrieved chunk count, validation decision, critic score and reliability gate result.",
                "These trace fields help explain why an answer was produced. A single baseline answer normally exposes retrieved chunks and final text but not planning, validation or criticism decisions.",
            ],
        ),
        (
            "Direct Lookup Boundary",
            [
                "The full hierarchy is not intended for every direct fact lookup. If the question asks for a single value and the first retrieved chunk contains it, a baseline route may be faster and sufficient.",
                "The hierarchy is intended for high reliability questions that require decomposition, multi source evidence, missing information recognition or relationship reasoning.",
            ],
        ),
    ],
    "thesis_doc_03_graph_reasoning_case_study.docx": [
        (
            "Entity Chain Appendix",
            [
                "The incident chain has six operational entities. Orion Service emits invoice events. Billing API accepts those events. Cache Layer stores normalized invoice records. Search Index reads from the Cache Layer. Audit Dashboard displays Search Index status. Compliance Review uses the dashboard as evidence.",
                "The relationship direction matters. Compliance Review depends on Audit Dashboard; Audit Dashboard depends on Search Index; Search Index depends on Cache Layer. Blaming the final review process ignores the upstream field mismatch.",
            ],
        ),
        (
            "Field Mapping Appendix",
            [
                "The schema migration used invoice_state as the new field name. The Cache Layer expected status_code as the searchable status field. This mismatch caused records to be stored without the searchable status value.",
                "Incomplete records then entered the Search Index. The Audit Dashboard showed pending review because it displayed indexed state rather than original Billing API payloads.",
            ],
        ),
        (
            "Ownership Appendix",
            [
                "Finance Systems owns the Billing API and is responsible for mapping invoice_state back to status_code during recovery. Reliability owns the Cache Layer and clears stale entries. Platform owns the Search Index and triggers index rebuilds.",
                "Compliance verifies the dashboard after corrected invoices are visible. Compliance does not own the Cache Layer, Search Index or schema migration.",
            ],
        ),
        (
            "Control Appendix",
            [
                "The compatibility contract is aimed at the boundary between Billing API and Cache Layer. It prevents future schema changes from breaking the expected searchable status field.",
                "The relationship test is aimed at the downstream path from Search Index to Audit Dashboard. It checks whether the expected status value survives migration and appears in the dashboard.",
            ],
        ),
        (
            "Missing Incident Appendix",
            [
                "The incident record deliberately omits customer names, exact revenue impact, affected invoice count, public severity grade and final regulatory conclusion.",
                "An answer that invents those missing details is not reliable. The correct answer should mark those details as unavailable in the provided document.",
            ],
        ),
    ],
    "thesis_doc_04_evaluation_and_ablation_report.docx": [
        (
            "Metric Interpretation Appendix",
            [
                "RAGAS gives a useful external view of faithfulness, answer relevance, context precision and context recall. It does not directly explain whether a system handled missing information, graph chains or validation retries correctly.",
                "The thesis therefore adds custom metrics for missing information accuracy, multi hop coverage, graph reasoning success, evidence support rate, citation rate, latency and retrieved chunk count.",
            ],
        ),
        (
            "Complex Benchmark Appendix",
            [
                "The complex benchmark focuses on multi hop relationship reasoning, graph dependency reasoning, evidence selection, missing information handling, self reflection and ablation explanation.",
                "It is not designed to prove that agentic RAG is better for every query. It is designed to test the type of reliability oriented tasks that motivate a hierarchical multi agent framework.",
            ],
        ),
        (
            "Ablation Detail Appendix",
            [
                "Removing BM25 affects exact lexical evidence because rare labels and numerical facts become harder to retrieve. Removing the Critic Agent affects faithfulness because weak answers are no longer reviewed after writing.",
                "Removing the Reliability Gate affects unsupported final answers because deterministic citation and context checks are removed. Removing graph search affects questions that connect services, teams, risks and downstream effects.",
            ],
        ),
        (
            "Latency Discussion Appendix",
            [
                "Latency is a limitation of the proposed system. The additional time comes from planning, decomposition, retrieval coordination, validation, writing, criticism and final gate checks.",
                "The report interprets the latency as acceptable for high reliability settings but unnecessary for low risk direct lookup settings.",
            ],
        ),
        (
            "Missing External Result Appendix",
            [
                "The report does not contain a public leaderboard rank, external competition score or production revenue result. It describes a controlled document benchmark rather than a public competition submission.",
                "Any answer claiming a leaderboard position or external score should be treated as unsupported by the provided document.",
            ],
        ),
    ],
}


def _sections_for_doc(spec: dict) -> list[tuple[str, list[str]]]:
    return list(spec["sections"]) + STRESS_APPENDIX_SECTIONS.get(spec["filename"], [])


def case(document, question, ground_truth, category, terms, *, graph=(), evidence=(), missing=False):
    return {
        "document": document,
        "question": question,
        "ground_truth": ground_truth,
        "category": category,
        "difficulty": "hard",
        "expected_missing": missing,
        "multi_hop_terms": list(terms),
        "graph_terms": list(graph),
        "evidence_terms": list(evidence or terms),
    }


COMPLEX_DATASET = [
    case("thesis_doc_01_rag_system_design.docx", "Explain how the RAG design connects retrieved chunks, citation rules and the reliability gate to prevent unsupported answers.", "Retrieved chunks provide evidence, citations connect claims to those chunks, and the reliability gate checks citations, available context and final answer support so unsupported fluent answers are not returned.", "self-reflection reliability reasoning", ["retrieved chunks", "citations", "reliability gate", "unsupported"], evidence=["citation", "reliability gate", "fallback"]),
    case("thesis_doc_01_rag_system_design.docx", "Why can a vector only baseline fail on broad summary questions, and how does the agentic design reduce that risk?", "A vector only baseline may return one narrow child chunk and miss evidence in other sections. The agentic design reduces the risk by combining vector and BM25 retrieval, synthesising duplicate evidence and checking the final answer.", "evaluation and ablation explanation", ["vector only", "narrow child chunk", "BM25", "synthesis", "final answer"], evidence=["failure", "BM25", "synthesising"]),
    case("thesis_doc_01_rag_system_design.docx", "If a rollback question mentions an approval exception that is not in the document, what should the system answer and why?", "The system should describe the available rollback rule if retrieved, then state that the approval exception is not provided because the missing information rule requires the writer to avoid unsupported claims.", "missing information handling", ["rollback", "approval exception", "not provided", "unsupported"], missing=True, evidence=["missing information", "unsupported", "provided documents"]),
    case("thesis_doc_01_rag_system_design.docx", "Compare the reliability value and latency cost of the agentic system against the baseline in this design.", "The agentic system improves reliability through coordinated retrieval, synthesis and final answer checks, but it is slower because it performs multiple retrieval and quality control steps. The baseline is faster because it uses one vector retrieval and one generation step.", "evaluation and ablation explanation", ["coordinated retrieval", "synthesis", "quality control", "latency", "baseline"], evidence=["tradeoff", "latency", "quality control"]),
    case("thesis_doc_01_rag_system_design.docx", "Which evidence supports using BM25 together with vector retrieval for reliable RAG?", "Vector retrieval handles paraphrased semantic questions, while BM25 handles exact terms, names, rare policy labels and measurements. The coordinator merges both so semantic and lexical evidence can be considered.", "evidence selection", ["vector retrieval", "BM25", "exact terms", "paraphrased", "coordinator"], evidence=["BM25", "exact terms", "coordinator"]),
    case("thesis_doc_01_rag_system_design.docx", "How should the fallback rule behave when the reliability gate fails but retrieved context is available?", "If the reliability gate fails while context is available, the system can fall back to a deterministic evidence summary instead of returning an unsupported fluent answer.", "self-reflection reliability reasoning", ["reliability gate", "fails", "context", "deterministic evidence summary", "unsupported"], evidence=["fallback", "deterministic evidence summary", "unsupported"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "Trace the full workflow for a complex query from planning to final reliability checking.", "A complex query flows through Planner, Query Decomposer, Retrieval Coordinator, Validator, Synthesis Agent, Writer Agent, Critic Agent and Reliability Gate.", "multi-hop relationship reasoning", ["Planner", "Query Decomposer", "Retrieval Coordinator", "Validator", "Synthesis Agent", "Writer Agent", "Critic Agent", "Reliability Gate"], graph=["planner", "retrieval coordinator", "critic agent", "reliability gate"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "How does validator retry make the agentic workflow more reliable than a single retrieval baseline?", "The Validator can detect insufficient evidence and return the workflow to the Retrieval Coordinator for another retrieval round, while a single retrieval baseline proceeds directly with the first retrieved evidence.", "self-reflection reliability reasoning", ["Validator", "insufficient", "Retrieval Coordinator", "another retrieval round", "baseline"], evidence=["validation", "retry", "insufficient"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "Explain the connection between the Writer Agent, Critic Agent and regeneration loop.", "The Writer Agent produces an answer, the Critic Agent evaluates accuracy, relevance, completeness and citation quality, and if the threshold is not satisfied the Critic returns feedback so the Writer can generate a revised answer.", "multi-hop relationship reasoning", ["Writer Agent", "Critic Agent", "accuracy", "citation quality", "feedback", "revised answer"], graph=["writer agent", "critic agent", "regeneration"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "Why does the benchmark disable critic regeneration while still recording critic scores?", "Critic regeneration can be disabled to keep experiments practical while critic scores are still recorded, preserving evidence of self reflection behaviour without excessive latency.", "evaluation and ablation explanation", ["critic regeneration", "disabled", "practical", "critic scores", "latency"], evidence=["benchmark", "critic scores", "practical"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "Which metadata fields make the multi agent workflow more transparent than black box generation?", "The workflow stores selected strategy, sub query count, retrieval agents used, retrieved chunk count, validation decision, critic score and final reliability status, making the process transparent.", "evidence selection", ["strategy", "sub query count", "retrieval agents", "validation decision", "critic score", "reliability status"], evidence=["metadata", "transparent", "critic score"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "When should the full hierarchy be preferred over baseline RAG according to the workflow document?", "The full hierarchy should be preferred for high reliability questions that need decomposition, multi source evidence, missing information handling or relationship reasoning, while direct facts may use baseline retrieval.", "evaluation and ablation explanation", ["high reliability", "decomposition", "multi source evidence", "missing information", "relationship reasoning", "direct facts"], evidence=["not intended", "baseline", "high reliability"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Trace the dependency path from Orion Service to Compliance Review in the incident.", "The dependency path is Orion Service to Billing API to Cache Layer to Search Index to Audit Dashboard to Compliance Review.", "graph dependency reasoning", ["Orion Service", "Billing API", "Cache Layer", "Search Index", "Audit Dashboard", "Compliance Review"], graph=["orion service", "billing api", "cache layer", "search index", "audit dashboard", "compliance review"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Explain why a Cache Layer failure can affect Compliance Review even when the dashboard is healthy.", "The Cache Layer feeds the Search Index, the Audit Dashboard displays invoice status from the Search Index, and Compliance Review depends on the dashboard, so Cache Layer failure can affect Compliance Review through the dependency chain.", "graph dependency reasoning", ["Cache Layer", "Search Index", "Audit Dashboard", "Compliance Review", "dependency chain"], graph=["cache layer", "search index", "audit dashboard", "compliance review"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Connect the schema migration field change to the downstream compliance effect.", "The migration replaced status_code with invoice_state. The Cache Layer still expected status_code, records lacked searchable status, the Search Index indexed incomplete records, the Audit Dashboard showed pending review, and Compliance Review was affected.", "graph dependency reasoning", ["invoice_state", "status_code", "Cache Layer", "Search Index", "Audit Dashboard", "Compliance Review"], graph=["invoice_state", "status_code", "cache layer", "search index", "audit dashboard"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Which evidence shows that the Audit Dashboard was only a downstream display component, not the root cause?", "The document states that the dashboard only displayed the indexed state, while the direct cause was the mismatch between invoice_state and status_code in the Cache Layer.", "evidence selection", ["dashboard", "indexed state", "direct cause", "invoice_state", "status_code", "Cache Layer"], evidence=["dashboard", "indexed state", "direct cause"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "How do the ownership responsibilities relate to the recovery procedure?", "Finance Systems owns the Billing API and maps invoice_state back to status_code, Reliability owns the Cache Layer and clears stale entries, Platform owns the Search Index and rebuilds it, and Compliance verifies the Audit Dashboard.", "multi-hop relationship reasoning", ["Finance Systems", "Billing API", "Reliability", "Cache Layer", "Platform", "Search Index", "Compliance"], graph=["finance systems", "billing api", "reliability", "cache layer", "platform", "search index", "compliance"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Why are a compatibility contract and relationship test recommended together?", "The compatibility contract controls the Billing API and Cache Layer schema boundary, while the relationship test checks whether the Search Index and Audit Dashboard receive the expected status value after migration.", "graph dependency reasoning", ["compatibility contract", "Billing API", "Cache Layer", "relationship test", "Search Index", "Audit Dashboard"], graph=["billing api", "cache layer", "search index", "audit dashboard"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "What exact revenue impact and final regulatory conclusion are reported for the Orion incident?", "The document does not provide exact revenue impact or a final regulatory conclusion.", "missing information handling", ["revenue impact", "regulatory conclusion", "not provide"], missing=True, evidence=["does not provide", "revenue impact", "regulatory conclusion"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "Explain how the main result summary supports the claim that agentic RAG improves reliability but costs latency.", "The proposed system improves RAGAS overall from 0.742 to 0.846, faithfulness from 0.710 to 0.890, citation rate from 0.78 to 0.96 and context usage from 0.46 to 0.61, while latency increases from 18.4 to 31.7 seconds.", "evaluation and ablation explanation", ["0.742", "0.846", "0.710", "0.890", "0.78", "0.96", "18.4", "31.7"], evidence=["faithfulness", "citation rate", "latency"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "Which ablation result best supports keeping the Reliability Gate in the proposed system?", "Removing the Reliability Gate increased unsupported final answers from 3 percent to 11 percent, supporting the need to keep the gate.", "evidence selection", ["Reliability Gate", "unsupported final answers", "3 percent", "11 percent"], evidence=["Reliability Gate", "unsupported", "11 percent"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "How do BM25, Critic Agent and Reliability Gate removals affect different parts of reliability?", "Removing BM25 reduces exact fact retrieval by 12 percent, removing the Critic Agent reduces faithfulness by 9 percent, and removing the Reliability Gate increases unsupported final answers from 3 percent to 11 percent.", "evaluation and ablation explanation", ["BM25", "12 percent", "Critic Agent", "9 percent", "Reliability Gate", "11 percent"], evidence=["BM25", "Critic Agent", "Reliability Gate"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "Why does graph search matter specifically for the graph reasoning subset?", "Graph search matters because relationship questions require connecting services, teams, risks and downstream effects. With graph search the subset scored 0.83 overall, while without graph search it scored 0.71.", "graph dependency reasoning", ["graph search", "relationship questions", "services", "teams", "risks", "downstream effects", "0.83", "0.71"], graph=["graph search", "relationship questions", "services", "teams", "risks"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "What public leaderboard rank and external competition score are reported for the proposed system?", "The document does not provide a public leaderboard rank or external competition score.", "missing information handling", ["leaderboard", "competition score", "not provide"], missing=True, evidence=["not provide", "leaderboard", "competition"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "How should the latency cost be interpreted in high reliability and low risk settings?", "The latency cost should be reported as a limitation. It is acceptable in high reliability settings requiring validation or multi hop reasoning, while a baseline route may be preferable for low risk direct lookup settings.", "evaluation and ablation explanation", ["latency", "limitation", "high reliability", "multi hop reasoning", "baseline", "direct lookup"], evidence=["latency", "limitation", "baseline route"]),
]

# Add balanced variants to reach 40 cases without simple direct facts.
COMPLEX_DATASET.extend([
    case("thesis_doc_01_rag_system_design.docx", "Why is a citation valid only when it directly supports the sentence, and how does this relate to the reliability gate?", "A citation is valid only when the cited chunk directly supports the sentence. The reliability gate then checks citations, citation numbers and retrieved context so unsupported claims are blocked.", "self-reflection reliability reasoning", ["directly supports", "reliability gate", "citation numbers", "retrieved context", "unsupported"], evidence=["citation", "directly supports", "gate"]),
    case("thesis_doc_01_rag_system_design.docx", "What should the system do if a user asks for an unapproved rollback exception and only related deployment rules are retrieved?", "It should answer only with supported deployment or rollback rules and state that an unapproved rollback exception is not contained in the provided documents.", "missing information handling", ["rollback", "exception", "not contained", "provided documents"], missing=True, evidence=["not contain", "unsupported", "provided documents"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "How do Planner strategy selection and Query Decomposer output influence retrieval behaviour?", "The Planner selects a strategy based on complexity. Complex questions use the Query Decomposer to create sub questions, and the Retrieval Coordinator retrieves evidence for those sub questions.", "multi-hop relationship reasoning", ["Planner", "strategy", "Query Decomposer", "sub questions", "Retrieval Coordinator"], graph=["planner", "query decomposer", "retrieval coordinator"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "What information is not provided about the exact model hardware used by the workflow agents?", "The document does not provide exact model hardware for the workflow agents.", "missing information handling", ["model hardware", "not provide", "workflow agents"], missing=True, evidence=["not provide", "hardware"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "How does the field mismatch move through records, indexing, dashboard status and compliance review?", "The mismatch between invoice_state and status_code causes records to lack searchable status, the Search Index indexes incomplete records, the Audit Dashboard shows pending review, and Compliance Review is affected.", "graph dependency reasoning", ["invoice_state", "status_code", "records", "Search Index", "Audit Dashboard", "Compliance Review"], graph=["invoice_state", "status_code", "search index", "audit dashboard", "compliance review"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Which unavailable incident details should the system refuse to invent?", "The system should refuse to invent customer names, exact revenue impact, affected invoice count, public incident severity grade and final regulatory conclusion because they are not provided.", "missing information handling", ["customer names", "revenue impact", "affected invoice", "severity grade", "regulatory conclusion", "not provided"], missing=True, evidence=["does not provide", "customer names", "regulatory conclusion"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "Why is the complex benchmark not evidence that agentic RAG should replace baseline RAG for every query?", "The report says baseline can be faster and sufficient for simple factual questions, while the proposed system is most useful for citations, comparison, missing information, relationship reasoning or quality control.", "evaluation and ablation explanation", ["baseline", "faster", "simple factual", "citations", "missing information", "relationship reasoning", "quality control"], evidence=["does not claim", "simple factual", "quality control"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "How do custom metrics extend RAGAS for this thesis benchmark?", "Custom metrics add citation rate, citation based context usage, missing information accuracy, multi hop coverage, graph reasoning success, answer length, latency and retrieved chunk count to RAGAS metrics.", "evaluation and ablation explanation", ["citation rate", "context usage", "missing information accuracy", "multi hop coverage", "graph reasoning success", "latency"], evidence=["custom metrics", "RAGAS", "graph reasoning success"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "Which evidence shows the workflow is intended for complex reliability tasks rather than every direct fact lookup?", "The document states that direct facts may use baseline retrieval, while the full hierarchy is intended for high reliability questions needing decomposition, multi source evidence, missing information handling or relationship reasoning.", "evidence selection", ["direct facts", "baseline retrieval", "full hierarchy", "high reliability", "decomposition", "relationship reasoning"], evidence=["direct facts", "baseline", "full hierarchy"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "How do the Validator, Critic Agent and Reliability Gate provide three different reliability checks?", "The Validator checks evidence sufficiency before writing, the Critic Agent reviews answer quality and can request regeneration, and the Reliability Gate performs deterministic final checks before the response is returned.", "self-reflection reliability reasoning", ["Validator", "evidence sufficiency", "Critic Agent", "regeneration", "Reliability Gate", "final checks"], evidence=["Validator", "Critic Agent", "Reliability Gate"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "Why is the Orion case suitable for graph based reasoning rather than isolated fact retrieval?", "The case requires connecting services, fields, teams, risks and downstream effects, including Orion Service, Billing API, Cache Layer, Search Index, Audit Dashboard and Compliance Review.", "graph dependency reasoning", ["services", "fields", "teams", "risks", "downstream effects", "Compliance Review"], graph=["orion service", "billing api", "cache layer", "search index", "audit dashboard", "compliance review"]),
    case("thesis_doc_01_rag_system_design.docx", "How do semantic evidence and exact lexical evidence work together in the retrieval coordinator?", "Vector retrieval supplies semantic evidence for paraphrased questions, BM25 supplies exact lexical evidence for rare terms and policy labels, and the retrieval coordinator merges both.", "multi-hop relationship reasoning", ["semantic evidence", "exact lexical evidence", "vector retrieval", "BM25", "retrieval coordinator"], graph=["vector retrieval", "bm25", "retrieval coordinator"]),
    case("thesis_doc_04_evaluation_and_ablation_report.docx", "Which result links graph search to downstream effects, teams and risks?", "The ablation study says graph search improved the graph reasoning subset from 0.71 to 0.83 because it helped questions connecting services, teams, risks and downstream effects.", "evidence selection", ["graph search", "0.71", "0.83", "services", "teams", "risks", "downstream effects"], evidence=["graph search", "0.83", "downstream effects"]),
    case("thesis_doc_02_multi_agent_workflow.docx", "How does recording validation decision and critic score support evaluation?", "Recording validation decision and critic score, along with strategy and retrieval metadata, makes the workflow transparent and allows evaluation of evidence sufficiency and answer quality control.", "self-reflection reliability reasoning", ["validation decision", "critic score", "strategy", "retrieval metadata", "transparent", "quality control"], evidence=["metadata", "validation decision", "critic score"]),
    case("thesis_doc_03_graph_reasoning_case_study.docx", "If someone blames Compliance Review for the incident, how should the graph chain correct that explanation?", "The graph chain shows Compliance Review depends on the Audit Dashboard, which depends on the Search Index, which depends on the Cache Layer. The direct cause was the Cache Layer field mismatch, not Compliance Review.", "graph dependency reasoning", ["Compliance Review", "Audit Dashboard", "Search Index", "Cache Layer", "direct cause", "field mismatch"], graph=["compliance review", "audit dashboard", "search index", "cache layer"]),
])


GRAPH_RELATIONSHIPS = [
    ("planner agent", "selects", "strategy"),
    ("query decomposer", "creates", "sub questions"),
    ("retrieval coordinator", "calls", "vector search agent"),
    ("retrieval coordinator", "calls", "keyword search agent"),
    ("retrieval coordinator", "calls", "graph search agent"),
    ("validator", "requests", "retrieval retry"),
    ("writer agent", "is_reviewed_by", "critic agent"),
    ("critic agent", "requests", "regeneration"),
    ("reliability gate", "checks", "citations"),
    ("orion service", "publishes_to", "billing api"),
    ("billing api", "writes_to", "cache layer"),
    ("cache layer", "feeds", "search index"),
    ("search index", "feeds", "audit dashboard"),
    ("audit dashboard", "supports", "compliance review"),
    ("schema migration", "introduced", "invoice_state"),
    ("cache layer", "expected", "status_code"),
    ("graph search", "improves", "relationship questions"),
    ("bm25", "improves", "exact fact retrieval"),
    ("critic agent", "improves", "faithfulness"),
    ("reliability gate", "reduces", "unsupported final answers"),
]


def _set_styles(doc: WordDocument) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Heading 1", 16, "1F4D78"),
        ("Heading 2", 13, "2E74B5"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def write_docx_files() -> None:
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    for spec in DOCS:
        doc = WordDocument()
        _set_styles(doc)
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(spec["title"])
        run.bold = True
        run.font.size = Pt(18)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor.from_string("1F4D78")
        for heading, paragraphs in _sections_for_doc(spec):
            doc.add_heading(heading, level=2)
            for text in paragraphs:
                p = doc.add_paragraph(text)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        path = CORPUS_DIR / spec["filename"]
        doc.core_properties.author = "Huang Xuan"
        doc.core_properties.title = spec["title"]
        doc.core_properties.subject = "Controlled complex thesis evaluation corpus"
        doc.save(path)
        print(f"Wrote {path}")


def write_complex_dataset() -> None:
    COMPLEX_DATASET_PATH.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "description": "Complex reasoning stress benchmark for hierarchical agentic RAG evaluation.",
        "documents": [doc["filename"] for doc in DOCS],
        "test_cases": COMPLEX_DATASET,
    }
    COMPLEX_DATASET_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Wrote {COMPLEX_DATASET_PATH}")


def _slug(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_").lower()
    return slug[:40] or "doc"


def _prefix_chunk_ids(parent_chunks, child_chunks, prefix: str) -> None:
    parent_id_map = {}
    for parent in parent_chunks:
        old_id = parent.chunk_id
        new_id = f"{prefix}_{old_id}"
        parent_id_map[old_id] = new_id
        parent.chunk_id = new_id
    for child in child_chunks:
        old_id = child.chunk_id
        child.chunk_id = f"{prefix}_{old_id}"
        if child.parent_id in parent_id_map:
            child.parent_id = parent_id_map[child.parent_id]
    for parent in parent_chunks:
        parent.children_ids = [f"{prefix}_{child_id}" for child_id in parent.children_ids]


def build_index() -> None:
    from dotenv import load_dotenv
    from src.ingestion.document_loader import DocumentLoader
    from src.ingestion.embedder import EmbeddingGenerator
    from src.ingestion.hierarchical_chunker import HierarchicalChunker
    from src.storage.chroma_store import ChromaVectorStore
    from src.retrieval.bm25_index import BM25Index

    load_dotenv(ROOT / ".env")
    if CHROMA_DIR.exists():
        shutil.rmtree(CHROMA_DIR)
    if BM25_PATH.exists():
        BM25_PATH.unlink()
    store = ChromaVectorStore(persist_directory=str(CHROMA_DIR))
    embedder = EmbeddingGenerator()
    loader = DocumentLoader()
    chunker = HierarchicalChunker(parent_size=700, child_size=180, child_overlap=35)
    for spec in DOCS:
        path = CORPUS_DIR / spec["filename"]
        loaded = loader.load(str(path))
        metadata = {"filename": spec["filename"], "source": "thesis_complex_corpus", "title": spec["title"]}
        parent_chunks, child_chunks = chunker.chunk_text(loaded.text, doc_id=loaded.doc_id, metadata=metadata)
        _prefix_chunk_ids(parent_chunks, child_chunks, _slug(path.stem))
        for chunk in parent_chunks + child_chunks:
            chunk.embedding = embedder.generate([chunk.text])[0]
        store.add_chunks(parent_chunks, child_chunks, filename=spec["filename"])
        print(f"Indexed {spec['filename']}: {len(parent_chunks)} parents, {len(child_chunks)} children")
    bm25 = BM25Index(index_path=str(BM25_PATH))
    bm25.build_from_vector_store(store)
    bm25.save()
    print(f"Wrote {BM25_PATH}")


def build_graph() -> None:
    from src.graph.graph_builder import KnowledgeGraph
    from src.graph.relationship_extractor import Relationship

    GRAPH_PATH.parent.mkdir(parents=True, exist_ok=True)
    kg = KnowledgeGraph()
    for source, relation, target in GRAPH_RELATIONSHIPS:
        kg.add_relationship(Relationship(source, relation, target, confidence=0.96))
    kg.save(str(GRAPH_PATH))
    print(f"Wrote {GRAPH_PATH}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-index", action="store_true")
    parser.add_argument("--no-graph", action="store_true")
    args = parser.parse_args()
    write_docx_files()
    write_complex_dataset()
    if not args.no_index:
        build_index()
    if not args.no_graph:
        build_graph()


if __name__ == "__main__":
    main()
